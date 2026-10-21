"""Coleta do myProfit no CRV e o status de fechamento que o CB lê.

Sem banco: autenticação da rota (o token do vínculo CB <-> CRV), a leitura do
pacote (tabela ou lista, datas, privacidade do formato) e os retratos de posição.
Com banco: a importação do pacote pelo comando (negócios, retratos por data,
proventos conferidos, reimportar sem duplicar), a conciliação do exterior e os
itens do status.
"""

from __future__ import annotations

import json
from datetime import date
from decimal import Decimal
from uuid import uuid4

import pytest
from sqlalchemy import func, select

from app import PUBLIC_ENDPOINTS
from app.executions import coleta
from app.executions.conciliacao import ONLY_CRV, QUANTITY, build_exterior, save_position_snapshot
from app.models import (
    Broker,
    Dividend,
    Execution,
    IncomeKind,
    Market,
    MyprofitColeta,
    MyprofitPosition,
    Portfolio,
    Position,
    Side,
    Ticker,
    User,
)
from tests.myprofit_sintetico import linha, relatorio_xlsx

D = Decimal
TOKEN = "t" * 40
COLUNAS = ["ticker", "corretora", "grupo", "classe", "quantidade", "preco_medio", "custo_brl", "valor_brl", "valor_usd"]


def _tsv_dos_negocios(linhas) -> str:
    """O histórico sintético como o navegador o extrai: TSV, titular anonimizado."""
    from io import BytesIO

    from openpyxl import load_workbook

    planilha = load_workbook(BytesIO(relatorio_xlsx(linhas)), read_only=True).active
    valores = [list(row) for row in planilha.iter_rows(values_only=True)]
    saida = ["Emitido por: TITULAR ANONIMIZADO - 000.000.000-00", str(valores[1][0])]
    for row in valores[2:]:
        saida.append("\t".join("" if c is None else str(c) for c in row))
    return "\n".join(saida)


def _pacote(sufixo="", *, negocios="", proventos=None, posicoes=None, ate="2026-10-20") -> bytes:
    posicoes = posicoes if posicoes is not None else {
        "2026-09-30": {"colunas": COLUNAS, "linhas": [
            [f"ACAO{sufixo}", "Genial", "Ações", "STOCK", 100, 10, 1000, 1100, 200],
            [f"EUA{sufixo}", "Avenue", "Ações EUA", "STOCK_USA", 50, 5, 250, 260, 50],
            ["5CYTYK9", "Avenue", "Renda Fixa EUA", "FIXED_INCOME_USA", 2000, 4.1, 8204, 8988, 1790],
        ]},
        "2026-10-20": {"colunas": COLUNAS, "linhas": [
            [f"ACAO{sufixo}", "Genial", "Ações", "STOCK", 100, 10, 1000, 1150, 210],
        ]},
    }
    return json.dumps({
        "formato": "myprofit-coleta/1", "de": "2026-09-01", "ate": ate, "posicoes": posicoes,
        "cambio": [], "proventos": proventos or [], "negocios_tsv": negocios,
    }).encode()


# --- sem banco ------------------------------------------------------------------

def test_rota_e_publica_e_autenticada_pelo_token_do_vinculo(client, app):
    assert "portfolio.fechamento_status" in PUBLIC_ENDPOINTS
    app.config["CB_PATRIMONIO_TOKEN"] = ""
    assert client.get("/fechamento/status?ate=2026-10-20").status_code == 503
    app.config["CB_PATRIMONIO_TOKEN"] = TOKEN
    assert client.get("/fechamento/status?ate=2026-10-20").status_code == 401
    assert client.get("/fechamento/status?ate=2026-10-20", headers={"Authorization": "Bearer outro"}).status_code == 401
    resposta = client.get("/fechamento/status?ate=ontem", headers={"Authorization": f"Bearer {TOKEN}"})
    assert resposta.status_code == 400


def test_pacote_aceita_tabela_e_lista_e_recusa_o_que_nao_e_coleta():
    pacote = coleta.ler(_pacote("X"))
    assert sorted(pacote.posicoes) == [date(2026, 9, 30), date(2026, 10, 20)]
    assert pacote.posicoes[date(2026, 9, 30)][0]["ticker"] == "ACAOX"
    lista = coleta.ler(_pacote(posicoes={"2026-10-20": [{"ticker": "A", "grupo": "Ações", "quantidade": 1}]}))
    assert lista.posicoes[date(2026, 10, 20)][0]["ticker"] == "A"
    with pytest.raises(coleta.ColetaError):
        coleta.ler(b'{"formato": "outro"}')
    with pytest.raises(coleta.ColetaError, match="depois da data"):
        coleta.ler(_pacote(posicoes={"2026-11-01": []}))
    with pytest.raises(coleta.ColetaError, match="formato"):
        coleta.ler(_pacote(posicoes={"2026-10-20": {"colunas": "x"}}))


def test_retratos_marcam_so_a_b3_como_escopo():
    retratos = coleta.retratos(coleta.ler(_pacote("X")))
    setembro = retratos[0]
    assert setembro.as_of == date(2026, 9, 30)
    assert [(row.asset, row.in_scope) for row in setembro.rows] == [
        ("ACAOX", True), ("EUAX", False), ("5CYTYK9", False),
    ]
    # Cada data tem seu próprio sha: reimportar uma data substitui só ela.
    assert retratos[0].sha256 != retratos[1].sha256


# --- com banco ------------------------------------------------------------------

@pytest.fixture
def cenario(sessao, app_com_banco, monkeypatch):
    sufixo = uuid4().hex[:6].upper()
    dono = User(username=f"fech-{sufixo.lower()}", password_hash="hash")
    sessao.add(dono)
    genial = sessao.scalar(select(Broker).where(Broker.name == "Genial"))
    if genial is None:
        genial = Broker(name="Genial", acronym=f"G{sufixo[:4]}")
        sessao.add(genial)
    avenue = Broker(name=f"Avenue {sufixo}", acronym=f"A{sufixo[:4]}")
    sessao.add(avenue)
    sessao.flush()
    real = Portfolio(name=f"Real {sufixo}", owner_ref=dono, currency="BRL")
    usd = Portfolio(name=f"USD {sufixo}", owner_ref=dono, currency="USD")
    acao = Ticker(symbol=f"ACAO{sufixo}", trading_name="Ação", market=Market.B3, rtd_market_code="B", currency="BRL")
    eua = Ticker(symbol=f"EUA{sufixo}", trading_name="EUA", market=Market.NYSE, rtd_market_code="N", currency="USD")
    sessao.add_all([real, usd, acao, eua])
    sessao.flush()
    sessao.add_all([
        Position(owner_id=dono.id, broker_id=genial.id, ticker_id=acao.id, portfolio_id=real.id,
                 quantity=D("100"), average_cost=D("10"), side=Side.BUY, opened_on=date(2026, 3, 2)),
        Position(owner_id=dono.id, broker_id=avenue.id, ticker_id=eua.id, portfolio_id=usd.id,
                 quantity=D("60"), average_cost=D("5"), side=Side.BUY, opened_on=date(2026, 3, 2)),
    ])
    sessao.flush()
    monkeypatch.setitem(app_com_banco.config, "MYPROFIT_INSTITUICOES", json.dumps({"GENIAL": "Genial"}))
    monkeypatch.setitem(app_com_banco.config, "CB_PATRIMONIO_TOKEN", TOKEN)
    monkeypatch.setitem(app_com_banco.config, "PATRIMONIO_OWNER_ID", str(dono.id))
    monkeypatch.setitem(app_com_banco.config, "CB_PATRIMONIO_URL", "")
    return {"dono": dono, "genial": genial, "acao": acao, "eua": eua, "sufixo": sufixo}


def _importar(c, conteudo: bytes):
    pacote = coleta.ler(conteudo)
    for retrato in coleta.retratos(pacote):
        save_position_snapshot(c["dono"].id, retrato)
    conferencia = coleta.conferir_proventos(c["dono"].id, pacote, date(2026, 1, 1))
    return coleta.registrar(c["dono"].id, pacote, conferencia), conferencia


@pytest.mark.banco
def test_exterior_compara_so_quantidade_na_data(sessao, cenario):
    c = cenario
    _importar(c, _pacote(c["sufixo"]))
    achados = {f.asset: f for f in build_exterior(c["dono"].id, date(2026, 9, 30))}
    eua = achados[f"EUA{c['sufixo']}"]
    assert (eua.state, eua.myprofit.quantity, eua.crv.quantity) == (QUANTITY, D("50"), D("60"))
    # A renda fixa EUA mora no CB: fica fora da comparação.
    assert "5CYTYK9" not in achados
    # Ativo da B3 não entra na conciliação do exterior.
    assert f"ACAO{c['sufixo']}" not in achados or achados[f"ACAO{c['sufixo']}"].state == ONLY_CRV


@pytest.mark.banco
def test_proventos_conferem_pelo_liquido_e_a_diferenca_aparece(sessao, cenario):
    c = cenario
    aluguel = Ticker(symbol=f"ALUG{c['sufixo']}", trading_name="Aluguel", market=Market.B3, rtd_market_code="B", currency="BRL")
    sessao.add(aluguel)
    sessao.flush()
    sessao.add_all([
        Dividend(
            owner_id=c["dono"].id, broker_id=c["genial"].id, ticker_id=c["acao"].id, kind=IncomeKind.JCP,
            amount=D("85.00"), payment_date=date(2026, 9, 10),
        ),
        # Aluguel vindo do CB: cada lado usa uma convenção, então fica fora da
        # conferência; mas ele marca até onde o CB já chegou (30/09).
        Dividend(
            owner_id=c["dono"].id, broker_id=c["genial"].id, ticker_id=aluguel.id, kind=IncomeKind.ALUGUEL,
            amount=D("82.70"), gross_amount=D("100.00"), withheld_amount=D("17.30"),
            payment_date=date(2026, 9, 30), origin="cb", origin_id=f"cb-{c['sufixo']}",
        ),
    ])
    sessao.flush()
    proventos = [
        # JCP: o CRV guarda o líquido (100 - 15 de imposto).
        {"ticker": f"ACAO{c['sufixo']}", "tipo": "JCP", "corretora": "Genial", "valor": 100, "imposto": 15,
         "valor_usd": 19, "data_pagamento": "2026-09-10"},
        {"ticker": "OUTRA3", "tipo": "DIVIDENDO", "corretora": "Genial", "valor": 40, "imposto": 0,
         "valor_usd": 8, "data_pagamento": "2026-09-12"},
        {"ticker": f"ALUG{c['sufixo']}", "tipo": "ALUGUEL", "corretora": "Genial", "valor": 100, "imposto": 0,
         "valor_usd": 19, "data_pagamento": "2026-09-30"},
        # Depois do último provento que o CB trouxe (30/09): o extrato ainda não
        # chegou, então fica de fora.
        {"ticker": "OUTRA3", "tipo": "DIVIDENDO", "corretora": "Genial", "valor": 40, "imposto": 0,
         "valor_usd": 8, "data_pagamento": "2026-10-05"},
        # Corretora que o CB não cobre não entra.
        {"ticker": "Aluguel", "tipo": "ALUGUEL", "corretora": "Avenue", "valor": 30, "imposto": 0,
         "valor_usd": 6, "data_pagamento": "2026-09-17"},
    ]
    registro, conferencia = _importar(c, _pacote(c["sufixo"], proventos=proventos))
    assert conferencia.fim == date(2026, 9, 30)
    assert conferencia.conferem == 1
    assert conferencia.divergencias == ["OUTRA3 09/2026: myProfit R$ 40.00, CRV R$ 0.00"]
    assert registro.summary["proventos"]["divergencias"] == conferencia.divergencias


@pytest.mark.banco
def test_reimportar_o_mesmo_pacote_nao_duplica(sessao, cenario):
    c = cenario
    conteudo = _pacote(c["sufixo"])
    _importar(c, conteudo)
    _importar(c, conteudo)
    assert sessao.scalar(select(func.count()).select_from(MyprofitColeta).where(MyprofitColeta.owner_id == c["dono"].id)) == 1
    linhas = sessao.scalar(
        select(func.count()).select_from(MyprofitPosition).where(
            MyprofitPosition.owner_id == c["dono"].id, MyprofitPosition.as_of == date(2026, 9, 30)
        )
    )
    assert linhas == 3


@pytest.mark.banco
def test_status_resume_o_crv_para_o_roteiro_do_cb(sessao, cenario, app_com_banco):
    c = cenario
    negocios = _tsv_dos_negocios([
        linha(date(2026, 10, 15), "NC-1", f"ACAO{c['sufixo']}", 10, "-101.50"),
    ])
    pacote = coleta.ler(_pacote(c["sufixo"], negocios=negocios))
    from app.executions.service import apply_import, plan_import

    apply_import(c["dono"].id, plan_import(c["dono"].id, coleta.negocios(pacote)), file_name="coleta.json")
    for retrato in coleta.retratos(pacote):
        save_position_snapshot(c["dono"].id, retrato)
    coleta.registrar(c["dono"].id, pacote, coleta.conferir_proventos(c["dono"].id, pacote, date(2026, 1, 1)))
    sessao.flush()
    assert sessao.scalar(select(func.count()).select_from(Execution).where(Execution.owner_id == c["dono"].id)) == 1

    resposta = app_com_banco.test_client().get(
        "/fechamento/status?ate=2026-10-20", headers={"Authorization": f"Bearer {TOKEN}"}
    )
    assert resposta.status_code == 200
    corpo = resposta.get_json()
    assert corpo["contrato"] == "crv-fechamento/1"
    itens = {item["chave"]: item for item in corpo["itens"]}
    assert (itens["coleta"]["estado"], itens["coleta"]["texto"]) == ("em_dia", "Importada até 20/10")
    assert itens["negocios"]["estado"] == "falta"
    assert itens["negocios"]["texto"] == "1 negócio para aplicar às posições"
    assert itens["exterior"]["estado"] == "pendente"  # 50 no myProfit, 60 no CRV, em 30/09
    assert itens["pregoes"]["estado"] == "espera"  # CB não configurado neste teste
    assert {item["etapa"] for item in corpo["itens"]} <= {"myprofit", "conferencias"}
    texto = json.dumps(corpo, ensure_ascii=False)
    assert "000.000.000-00" not in texto
