from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.validation import parse_finite_decimal
from app.models import IncomeKind

REQUIRED_COLUMNS = (
    "Ativo",
    "Corretora",
    "Recebido",
    "Tipo",
    "Data Pgto.",
    "Data Com",
    "YOC",
    "DY",
    "Cotas",
    "Total investido",
    "Total atual",
    "Preço Médio",
    "Cotação",
    "Valor por Cota",
)
MAX_WORKBOOK_BYTES = 5 * 1024 * 1024


class DividendWorkbookError(ValueError):
    """Planilha que não pode ser importada com segurança."""


@dataclass(frozen=True, slots=True)
class DividendImportRow:
    source_row: int
    ticker: str
    broker: str
    amount: Decimal
    kind: IncomeKind
    payment_date: date
    com_date: date
    yield_on_cost: Decimal
    dividend_yield: Decimal
    quotas: Decimal
    invested_total: Decimal
    market_total: Decimal
    average_price: Decimal
    quoted_price: Decimal
    amount_per_share: Decimal | None


@dataclass(frozen=True, slots=True)
class DividendWorkbook:
    rows: list[DividendImportRow]
    rejected_rows: list[str]


def _normalized_header(value: Any) -> str:
    return " ".join(str(value or "").split()).casefold()


def _required_header_map(values: tuple[Any, ...]) -> dict[str, int] | None:
    indexes = {_normalized_header(value): index for index, value in enumerate(values)}
    required = {_normalized_header(column): column for column in REQUIRED_COLUMNS}
    if not required.keys() <= indexes.keys():
        return None
    return {column: indexes[_normalized_header(column)] for column in REQUIRED_COLUMNS}


def _as_date(value: Any, *, field_name: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        raw = value.strip()
        for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(raw, pattern).date()
            except ValueError:
                pass
    raise ValueError(f"Informe {field_name} válida.")


def _as_decimal(value: Any, *, field_name: str) -> Decimal:
    if value is None or isinstance(value, bool) or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"Informe {field_name}.")
    return parse_finite_decimal(str(value).strip(), field_name=field_name)


def _as_text(value: Any, *, field_name: str) -> str:
    if value is None:
        raise ValueError(f"Informe {field_name}.")
    result = str(value).strip()
    if not result:
        raise ValueError(f"Informe {field_name}.")
    return result


def _income_kind(value: Any) -> IncomeKind:
    raw = _as_text(value, field_name="o tipo").casefold()
    if raw == "reembolso":
        return IncomeKind.DIVIDENDO
    try:
        return IncomeKind(raw)
    except ValueError as exc:
        raise ValueError("Informe um tipo válido: Dividendo, JCP, Aluguel ou Reembolso.") from exc


def _parse_row(values: tuple[Any, ...], headers: dict[str, int], source_row: int) -> DividendImportRow:
    def cell(column: str) -> Any:
        return values[headers[column]] if headers[column] < len(values) else None

    amount = _as_decimal(cell("Recebido"), field_name="o valor recebido")
    # YOC e DY existem no arquivo como uma conferência visível para quem o
    # exportou. Não são a fonte do cálculo: os totais da linha prevalecem e
    # evitam que uma fórmula antiga da planilha seja perpetuada no CRV.
    source_yoc = _as_decimal(cell("YOC"), field_name="o YOC")
    source_dy = _as_decimal(cell("DY"), field_name="o DY")
    quotas = _as_decimal(cell("Cotas"), field_name="as cotas")
    invested_total = _as_decimal(cell("Total investido"), field_name="o total investido")
    market_total = _as_decimal(cell("Total atual"), field_name="o total atual")
    average_price = _as_decimal(cell("Preço Médio"), field_name="o preço médio")
    quoted_price = _as_decimal(cell("Cotação"), field_name="a cotação")
    _as_decimal(cell("Valor por Cota"), field_name="o valor por cota")
    payment_date = _as_date(cell("Data Pgto."), field_name="a data de pagamento")
    if amount <= 0:
        raise ValueError("O valor recebido deve ser positivo.")
    if payment_date > date.today():
        raise ValueError("A data de pagamento não pode estar no futuro.")
    financial_values = (
        source_yoc,
        source_dy,
        quotas,
        invested_total,
        market_total,
        average_price,
        quoted_price,
    )
    if any(value < 0 for value in financial_values):
        raise ValueError("YOC, DY, cotas e valores financeiros não podem ser negativos.")
    # A planilha registra taxas em pontos percentuais (0,42 significa 0,42%).
    # Mantemos essa unidade nos campos históricos para não misturar o dado de
    # origem com as taxas consolidadas, que são razões e recebem o filtro
    # ``percent`` na apresentação.
    yield_on_cost = (amount / invested_total * Decimal("100")) if invested_total else None
    dividend_yield = (amount / market_total * Decimal("100")) if market_total else None
    return DividendImportRow(
        source_row=source_row,
        ticker=_as_text(cell("Ativo"), field_name="o ativo").upper(),
        broker=_as_text(cell("Corretora"), field_name="a corretora"),
        amount=amount,
        kind=_income_kind(cell("Tipo")),
        payment_date=payment_date,
        com_date=_as_date(cell("Data Com"), field_name="a data com"),
        yield_on_cost=yield_on_cost,
        dividend_yield=dividend_yield,
        quotas=quotas,
        invested_total=invested_total,
        market_total=market_total,
        average_price=average_price,
        quoted_price=quoted_price,
        amount_per_share=(amount / quotas) if quotas else None,
    )


def read_dividend_workbook(content: bytes) -> DividendWorkbook:
    """Lê uma planilha XLSX, localizando a aba pelos cabeçalhos obrigatórios."""
    if not content:
        raise DividendWorkbookError("Selecione uma planilha XLSX.")
    if len(content) > MAX_WORKBOOK_BYTES:
        raise DividendWorkbookError("A planilha excede o limite de 5 MB.")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise DividendWorkbookError("Não foi possível ler a planilha XLSX.") from exc

    closest_headers: tuple[Any, ...] = ()
    closest_matches = -1
    required_headers = {_normalized_header(column) for column in REQUIRED_COLUMNS}
    for worksheet in workbook.worksheets:
        for source_row, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
            matches = len(required_headers & {_normalized_header(value) for value in values})
            if matches > closest_matches:
                closest_headers = values
                closest_matches = matches
            headers = _required_header_map(values)
            if headers is None:
                continue
            rows: list[DividendImportRow] = []
            rejected: list[str] = []
            for row_number, row_values in enumerate(
                worksheet.iter_rows(min_row=source_row + 1, values_only=True), start=source_row + 1
            ):
                if not any(value is not None and str(value).strip() for value in row_values):
                    continue
                try:
                    rows.append(_parse_row(row_values, headers, row_number))
                except ValueError as exc:
                    rejected.append(f"Linha {row_number}: {exc}")
            return DividendWorkbook(rows=rows, rejected_rows=rejected)

    present = {_normalized_header(value) for value in closest_headers}
    missing = [column for column in REQUIRED_COLUMNS if _normalized_header(column) not in present]
    raise DividendWorkbookError("Faltam as colunas obrigatórias: " + ", ".join(missing) + ".")
