"""Leitura do relatório "Histórico de investimentos" do myProfit (Excel).

O arquivo começa com duas linhas que identificam o titular: ``Emitido por:
<nome> - <CPF>`` e ``Período: dd/mm/aaaa até dd/mm/aaaa``. Do período se lê
apenas as duas datas; o nome e o CPF **nunca** saem desta função: não vão para
o resultado, para mensagem de erro nem para log. A terceira linha é o
cabeçalho da tabela.

O módulo é puro (sem Flask, sem ORM): recebe bytes e devolve linhas validadas.
Decidir a corretora, comparar com o que já está gravado e gravar é do
``app.executions.service``.
"""

from __future__ import annotations

import hashlib
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_EVEN, Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.validation import parse_finite_decimal

REQUIRED_COLUMNS = (
    "Data de negociação",
    "Instituição",
    "Número",
    "Documento",
    "Moeda",
    "Total de Taxas",
    "Ativo",
    "Grupo",
    "Quantidade",
    "Operação",
    "Tipo",
    "Preço sem taxas",
    "Preço com taxas",
    "Total sem taxas",
    "Total com taxas",
    "Obs",
)
MAX_WORKBOOK_BYTES = 2 * 1024 * 1024

#: O Excel conta dias a partir de 30/12/1899 (a época já embute o erro do ano
#: 1900 bissexto): 46028 é 06/01/2026.
_EXCEL_EPOCH = date(1899, 12, 30)
_DECIMAL_STEP = Decimal("0.00000001")
_DECIMAL_LIMIT = Decimal(10) ** 15  # Numeric(24, 8)

#: Grupo do relatório -> código gravado. Um grupo novo é recusado de propósito:
#: melhor parar e decidir do que gravar sem saber se pertence à B3.
ASSET_GROUPS = {
    "acoes": "acao",
    "opcoes": "opcao",
    "etf": "etf",
    "indice futuro": "indice_futuro",
    "tesouro direto": "tesouro_direto",
    "acoes eua": "acao_eua",
    "etf usa": "etf_usa",
}
#: Os grupos que a conferência e o uso futuro consideram.
B3_GROUPS = frozenset({"acao", "opcao", "etf"})

_PERIOD = re.compile(
    r"per[ií]odo:\s*(\d{2}/\d{2}/\d{4})\s*at[eé]\s*(\d{2}/\d{2}/\d{4})", re.IGNORECASE
)
_CPF_IN_NAME = re.compile(r"\d{3}\.?\d{3}\.?\d{3}-?\d{2}")


class ExecutionReportError(ValueError):
    """Arquivo que não pode ser importado com segurança.

    As mensagens nomeiam colunas e limites, nunca o conteúdo do cabeçalho.
    """


@dataclass(frozen=True, slots=True)
class ExecutionRow:
    source_row: int
    trade_date: date
    institution: str
    document: str
    document_type: str | None
    line_order: int
    asset: str
    asset_group: str
    side: str
    cash_operation: str
    quantity: Decimal
    price_without_fees: Decimal
    price_with_fees: Decimal
    fees_total: Decimal
    total_without_fees: Decimal
    total_with_fees: Decimal
    day_trade: bool
    currency: str
    in_scope: bool
    content_hash: str


@dataclass(frozen=True, slots=True)
class ExecutionReport:
    period_start: date
    period_end: date
    sha256: str
    rows: list[ExecutionRow]
    rejected_rows: list[str]


def _fold(value: Any) -> str:
    text = unicodedata.normalize("NFKD", " ".join(str(value or "").split()))
    return "".join(char for char in text if not unicodedata.combining(char)).casefold()


def is_in_scope(asset_group: str, document: str) -> bool:
    """Ações, opções e ETF da B3 — sem a taxa de aluguel de ações.

    O documento ``B3_ALUGUEL_*`` lança tarifas de centavos do aluguel como
    "compra" de uma unidade; não é um pregão e o CB não as liquida como tal.
    """
    return asset_group in B3_GROUPS and "ALUGUEL" not in document.upper()


def _as_date(value: Any, *, field_name: str) -> date:
    if isinstance(value, bool):
        raise ValueError(f"Informe {field_name} válida.")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float, Decimal)):
        serial = Decimal(str(value))
        if serial != serial.to_integral_value() or not 1 <= serial <= 2_958_465:
            raise ValueError(f"Informe {field_name} válida.")
        return _EXCEL_EPOCH + timedelta(days=int(serial))
    if isinstance(value, str):
        raw = value.strip()
        if raw.isdigit():
            return _as_date(int(raw), field_name=field_name)
        for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(raw, pattern).date()
            except ValueError:
                pass
    raise ValueError(f"Informe {field_name} válida.")


def _as_decimal(value: Any, *, field_name: str) -> Decimal:
    if value is None or isinstance(value, bool) or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"Informe {field_name}.")
    raw = str(value).strip()
    if isinstance(value, str) and "," in raw:
        # "1.234,56" -> "1234.56"; "12,5" -> "12.5".
        raw = raw.replace(".", "").replace(",", ".")
    try:
        parsed = parse_finite_decimal(raw, field_name=field_name)
        # Oito casas é o que a coluna guarda; arredondar aqui mantém o valor
        # lido igual ao gravado, e é o que impede uma linha repetida de parecer
        # alterada.
        result = parsed.quantize(_DECIMAL_STEP, rounding=ROUND_HALF_EVEN)
    except (ValueError, InvalidOperation) as exc:
        raise ValueError(f"Informe {field_name} válido.") from exc
    if abs(result) >= _DECIMAL_LIMIT:
        raise ValueError(f"Informe {field_name} válido.")
    return result


def _as_text(value: Any, *, field_name: str, limit: int) -> str:
    result = " ".join(str(value if value is not None else "").split())
    if not result:
        raise ValueError(f"Informe {field_name}.")
    if len(result) > limit:
        raise ValueError(f"{field_name.capitalize()} excede {limit} caracteres.")
    return result


def _canonical(value: Decimal) -> str:
    return format(value.normalize(), "f") if value else "0"


def _content_hash(row: dict[str, Any]) -> str:
    material = "|".join(
        [
            row["trade_date"].isoformat(),
            row["institution"].casefold(),
            row["document"],
            row["document_type"] or "",
            row["asset"],
            row["asset_group"],
            row["side"],
            row["cash_operation"],
            _canonical(row["quantity"]),
            _canonical(row["price_without_fees"]),
            _canonical(row["price_with_fees"]),
            _canonical(row["fees_total"]),
            _canonical(row["total_without_fees"]),
            _canonical(row["total_with_fees"]),
            "1" if row["day_trade"] else "0",
            row["currency"],
        ]
    )
    return hashlib.sha256(material.encode("utf-8")).hexdigest()


def _header_map(values: tuple[Any, ...]) -> dict[str, int] | None:
    indexes = {_fold(value): index for index, value in enumerate(values) if value is not None}
    wanted = {_fold(column): column for column in REQUIRED_COLUMNS}
    if not wanted.keys() <= indexes.keys():
        return None
    return {column: indexes[folded] for folded, column in wanted.items()}


def _parse_row(
    values: tuple[Any, ...], headers: dict[str, int], source_row: int, line_order: int
) -> ExecutionRow:
    def cell(column: str) -> Any:
        index = headers[column]
        return values[index] if index < len(values) else None

    group_raw = _fold(cell("Grupo"))
    asset_group = ASSET_GROUPS.get(group_raw)
    if asset_group is None:
        raise ValueError("o grupo do ativo não é reconhecido.")
    side_raw = _fold(cell("Tipo"))
    side = {"compra": "C", "venda": "V"}.get(side_raw)
    if side is None:
        raise ValueError("o tipo deve ser Compra ou Venda.")
    operation_raw = _fold(cell("Operação"))
    cash_operation = {"credito": "C", "debito": "D"}.get(operation_raw)
    if cash_operation is None:
        raise ValueError("a operação deve ser Crédito ou Débito.")
    currency = _as_text(cell("Moeda"), field_name="a moeda", limit=3).upper()
    if not re.fullmatch(r"[A-Z]{3}", currency):
        raise ValueError("a moeda deve ter três letras.")
    document_type_raw = " ".join(str(cell("Documento") or "").split())
    if len(document_type_raw) > 40:
        raise ValueError("o documento excede 40 caracteres.")
    obs = _fold(cell("Obs"))

    fields: dict[str, Any] = {
        "trade_date": _as_date(cell("Data de negociação"), field_name="a data de negociação"),
        "institution": _as_text(cell("Instituição"), field_name="a instituição", limit=40),
        "document": _as_text(cell("Número"), field_name="o número do documento", limit=80),
        "document_type": document_type_raw or None,
        "asset": _as_text(cell("Ativo"), field_name="o ativo", limit=80),
        "asset_group": asset_group,
        "side": side,
        "cash_operation": cash_operation,
        "quantity": _as_decimal(cell("Quantidade"), field_name="a quantidade"),
        "price_without_fees": _as_decimal(cell("Preço sem taxas"), field_name="o preço sem taxas"),
        "price_with_fees": _as_decimal(cell("Preço com taxas"), field_name="o preço com taxas"),
        "fees_total": _as_decimal(cell("Total de Taxas"), field_name="o total de taxas"),
        "total_without_fees": _as_decimal(cell("Total sem taxas"), field_name="o total sem taxas"),
        "total_with_fees": _as_decimal(cell("Total com taxas"), field_name="o total com taxas"),
        "day_trade": "day" in obs and "trade" in obs,
        "currency": currency,
    }
    if fields["quantity"] == 0:
        raise ValueError("a quantidade não pode ser zero.")
    if min(fields["price_without_fees"], fields["price_with_fees"], fields["fees_total"]) < 0:
        raise ValueError("preços e taxas não podem ser negativos.")
    return ExecutionRow(
        source_row=source_row,
        line_order=line_order,
        in_scope=is_in_scope(asset_group, fields["document"]),
        content_hash=_content_hash(fields),
        **fields,
    )


def safe_file_name(name: str | None) -> str | None:
    """Nome do arquivo para o registro do lote, sem sequência que pareça CPF.

    O nome do arquivo é escolhido por quem exporta e pode carregar dados
    pessoais; o que parece CPF sai, e o resto é só o nome da base.
    """
    if not name:
        return None
    base = name.replace("\\", "/").rsplit("/", 1)[-1].strip()
    base = _CPF_IN_NAME.sub("", base)
    return base[:120] or None


def load_sheets(content: bytes) -> list[list[tuple[Any, ...]]]:
    """Linhas de cada planilha do relatório, de um XLSX ou de um TSV.

    O TSV é o texto que o navegador embutido extrai do relatório (o Excel que o
    site oferece não sai da página): células separadas por tabulação, primeira
    linha ``Emitido por: ...`` já sem o titular. Célula vazia vira ``None``,
    como no Excel; números e datas chegam como texto e passam pelas mesmas
    conversões das células do XLSX.
    """
    if not content:
        raise ExecutionReportError("Selecione o relatório em formato XLSX ou TSV.")
    if len(content) > MAX_WORKBOOK_BYTES:
        raise ExecutionReportError("O relatório excede o limite de 2 MB.")
    if not content.startswith(b"PK"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ExecutionReportError("Não foi possível ler o relatório TSV (use UTF-8).") from exc
        lines = [
            tuple(cell if cell.strip() else None for cell in line.split("	"))
            for line in text.splitlines()
        ]
        return [lines]
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        return [list(worksheet.iter_rows(values_only=True)) for worksheet in workbook.worksheets]
    except (InvalidFileException, zipfile.BadZipFile, OSError, ValueError, KeyError) as exc:
        raise ExecutionReportError("Não foi possível ler o relatório XLSX.") from exc


def read_execution_report(content: bytes) -> ExecutionReport:
    """Lê o relatório e devolve período, linhas válidas e linhas recusadas.

    Falha com ``ExecutionReportError`` se o arquivo for vazio, grande demais,
    ilegível ou não tiver o cabeçalho e as colunas do relatório.
    """
    sheets = load_sheets(content)

    last_problem = (
        "O relatório não começa com as linhas 'Emitido por' e 'Período'; "
        "use o 'Histórico de investimentos' do myProfit."
    )
    for sheet_rows in sheets:
        preview = sheet_rows[:3]
        if len(preview) < 3:
            continue
        # A primeira linha só é conferida, nunca lida adiante: tem nome e CPF.
        first_cell = _fold(preview[0][0] if preview[0] else None)
        period = _PERIOD.search(str(preview[1][0] if preview[1] else ""))
        if not first_cell.startswith("emitido por") or period is None:
            last_problem = (
                "O relatório não começa com as linhas 'Emitido por' e 'Período'; "
                "use o 'Histórico de investimentos' do myProfit."
            )
            continue
        headers = _header_map(preview[2])
        if headers is None:
            present = {_fold(value) for value in preview[2] if value is not None}
            missing = [column for column in REQUIRED_COLUMNS if _fold(column) not in present]
            last_problem = "Faltam as colunas obrigatórias: " + ", ".join(missing) + "."
            continue
        try:
            period_start = _as_date(period.group(1), field_name="o início do período")
            period_end = _as_date(period.group(2), field_name="o fim do período")
        except ValueError as exc:
            raise ExecutionReportError("O período do relatório é inválido.") from exc
        if period_end < period_start:
            raise ExecutionReportError("O período do relatório é inválido.")

        rows: list[ExecutionRow] = []
        rejected: list[str] = []
        orders: dict[str, int] = {}
        for source_row, values in enumerate(sheet_rows[3:], start=4):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            document = " ".join(str(values[headers["Número"]] if headers["Número"] < len(values) else "").split())
            # A ordem conta toda linha do documento, inclusive a recusada, para
            # que consertar uma linha não desloque as vizinhas.
            orders[document] = orders.get(document, 0) + 1
            try:
                rows.append(_parse_row(values, headers, source_row, orders[document]))
            except ValueError as exc:
                rejected.append(f"Linha {source_row}: {exc}")
        return ExecutionReport(
            period_start=period_start,
            period_end=period_end,
            sha256=hashlib.sha256(content).hexdigest(),
            rows=rows,
            rejected_rows=rejected,
        )
    raise ExecutionReportError(last_problem)
